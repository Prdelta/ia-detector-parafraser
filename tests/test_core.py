import io
import os

os.environ["IADECCION_BINOCULARS"] = "0"  # pruebas rápidas sin descargar modelos
os.environ["IADECCION_PARAFRASEADOR"] = "0"
os.environ["IADECCION_CLASIFICADOR"] = "0"

import docx  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from backend.app.analyzer import analyze_text  # noqa: E402
from backend.app.detectors import meta, stylometry  # noqa: E402
from backend.app.extract import UnsupportedFile, extract_text  # noqa: E402
from backend.app.main import app  # noqa: E402
from backend.app.rewrite import guidance  # noqa: E402
from backend.app.segment import split_sentences  # noqa: E402

HUMANO = (
    "Ayer fui al mercado. Llovía a cántaros y, claro, se me olvidó el paraguas en el bus, como siempre. "
    "La señora de las frutas me regaló dos mangos porque dice que estoy flaco. ¿Flaco yo? "
    "Bueno, quizá un poco. Después pasé por la biblioteca de la facultad para devolver el libro de Vargas Llosa "
    "que debía desde marzo; la multa fue ridícula, menos de lo que cuesta un café.\n\n"
    "En la tarde intenté escribir el informe de laboratorio, pero el Excel se cerró tres veces y perdí la tabla. "
    "Terminé a las once. No quedó perfecto. Mañana lo reviso con Camila, que entiende mejor la parte de titulación."
)

IA = (
    "En la era digital, la educación desempeña un papel fundamental en el desarrollo de la sociedad. "
    "Además, es importante destacar que la tecnología ofrece una amplia gama de oportunidades para los estudiantes. "
    "Asimismo, cabe destacar que los docentes deben adoptar un enfoque integral para fomentar el aprendizaje. "
    "Por otro lado, la innovación educativa permite abordar los desafíos de manera efectiva y eficiente.\n\n"
    "En este sentido, la colaboración entre instituciones juega un papel crucial en la mejora de la calidad educativa. "
    "Además, es fundamental considerar los desafíos y oportunidades que surgen en el panorama actual. "
    "Por lo tanto, resulta esencial promover estrategias que potencien las competencias digitales de los estudiantes. "
    "En conclusión, la educación en la era digital representa un aspecto clave para el futuro de la sociedad."
)


def test_split_respeta_abreviaturas_y_decimales():
    s = split_sentences("El Dr. Pérez midió 3.5 cm. Luego habló con la Sra. Gómez. ¿Listo?")
    assert [x.text for x in s] == ["El Dr. Pérez midió 3.5 cm.", "Luego habló con la Sra. Gómez.", "¿Listo?"]


def test_posiciones_de_oraciones():
    text = "Primera oración.\n\nSegunda oración aquí."
    for s in split_sentences(text):
        assert text[s.start:s.end] == s.text
    assert [s.paragraph for s in split_sentences(text)] == [0, 1]


def test_estilometria_distingue_ejemplos_claros():
    h = stylometry.analyze(HUMANO, split_sentences(HUMANO))
    a = stylometry.analyze(IA, split_sentences(IA))
    assert a.score > 0.65 > 0.35 > h.score
    assert a.reasons


def test_texto_corto_rechazado():
    with pytest.raises(ValueError):
        analyze_text("Muy corto.")


def test_informe_completo():
    r = analyze_text(IA + "\n\n" + HUMANO)
    assert 0 <= r["probabilidad_ia"] <= 1
    assert r["confianza"] == "baja"  # sin modelo principal
    assert len(r["oraciones"]) > 10
    assert all(o["nivel"] in {"bajo", "medio", "alto"} for o in r["oraciones"])


def test_extraer_docx():
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("Hola mundo.")
    d.add_paragraph("Segundo párrafo.")
    d.save(buf)
    assert extract_text("a.docx", buf.getvalue()) == "Hola mundo.\n\nSegundo párrafo."
    with pytest.raises(UnsupportedFile):
        extract_text("a.exe", b"x")


def test_docx_con_citas_de_zotero_y_sin_bibliografia():
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    p = d.add_paragraph()
    # Cita insertada por un gestor de referencias: un control de contenido (w:sdt) dentro del párrafo.
    p._p.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:r><w:t>(Kratzert et al., 2019)</w:t></w:r>'
        '</w:sdtContent></w:sdt>'))
    p.add_run(", en Estados Unidos, entrenaron una red LSTM.")
    d.add_paragraph("Kratzert, F., Klotz, D., y Nearing, G. (2019). Towards learning universal hydrological "
                    "behaviors. Hydrology and Earth System Sciences, 23, 5089-5110.")
    tabla = d.add_table(rows=1, cols=1)
    tabla.cell(0, 0).text = "celda de tabla"
    buf = io.BytesIO()
    d.save(buf)
    assert extract_text("a.docx", buf.getvalue()) == (
        "(Kratzert et al., 2019), en Estados Unidos, entrenaron una red LSTM.")


def test_api():
    with TestClient(app) as client:
        assert client.get("/api/salud").json()["estado"] == "ok"
        assert client.get("/").status_code == 200
        r = client.post("/api/analizar", json={"texto": IA})
        assert r.status_code == 200 and "veredicto" in r.json()
        assert client.post("/api/analizar", json={"texto": "corto"}).status_code == 422
        r = client.post("/api/analizar-archivo", files={"archivo": ("t.txt", HUMANO.encode(), "text/plain")})
        assert r.status_code == 200 and r.json()["archivo"] == "t.txt"


def test_parrafos_con_guia():
    r = analyze_text(IA + "\n\n" + HUMANO)
    assert r["texto"].startswith("En la era digital")
    assert [p["indice"] for p in r["parrafos"]] == [0, 1, 2, 3]
    assert "\n\n".join(p["texto"] for p in r["parrafos"]) == r["texto"]
    primero = r["parrafos"][0]
    assert any("Frases hechas" in m for m in primero["motivos"])
    assert any("conectores" in m for m in primero["motivos"])
    assert primero["sugerencias"]
    assert r["parafrasis_disponible"] is False


def test_guia_parrafo_humano_sin_motivos():
    sents = split_sentences(HUMANO.split("\n\n")[0])
    g = guidance(sents, [[] for _ in sents], probability=0.1)
    assert g == {"motivos": [], "sugerencias": []}


def test_parafrasear_desactivado():
    with TestClient(app) as client:
        r = client.post("/api/parafrasear", json={"texto": "Un párrafo cualquiera con varias palabras."})
        assert r.status_code == 503


def test_meta_combina_senales():
    m = meta.Meta(coef=[-1.0, 0.5, 1.0], intercept=0.0, mean=[0.9, 0.5, 0.0], std=[0.1, 0.2, 2.0])
    ia = m.probability(meta.features(0.7, 0.8, 0.99))
    humano = m.probability(meta.features(1.05, 0.2, 0.01))
    assert ia > 0.9 and humano < 0.1
    assert m.probability(meta.features(0.9, 0.5, 0.5)) == 0.5


IA_EN = (
    "In today's rapidly evolving world, education plays a crucial role in shaping the future of society. "
    "Furthermore, it is important to note that technology offers a wide range of opportunities for students. "
    "Moreover, teachers must navigate the complexities of digital learning in a seamless way. "
    "Additionally, schools should foster a culture of innovation and collaboration among learners. "
    "In conclusion, the realm of education is not only changing quickly but also creating new challenges and opportunities. "
) * 4


def test_ingles_detectado_con_lexico_propio():
    r = analyze_text(IA_EN)
    assert r["idioma"] == "en"
    assert any("inglés" in a for a in r["avisos"])
    exprs = {e.lower() for o in r["oraciones"] for e in o["expresiones"]}
    assert {"plays a crucial role", "it is important to note", "in conclusion"} <= exprs
    assert r["senales"]["estilometria"]["rasgos"]["conectores_inicio"] > 0.5


def test_espanol_no_cambia_de_idioma():
    assert analyze_text(IA)["idioma"] == "es"
    assert analyze_text(HUMANO + " " + HUMANO)["idioma"] == "es"


class _ClasificadorInflado:
    model_id = "falso"

    def probability(self, text):
        return 0.99


def test_motivo_del_clasificador_solo_si_el_resultado_es_ia(monkeypatch):
    from backend.app import analyzer

    monkeypatch.setattr(analyzer.classifier, "get_classifier", lambda: _ClasificadorInflado())
    r = analyze_text(HUMANO + " " + HUMANO)
    assert r["probabilidad_ia"] < 0.65
    assert not any("clasificador entrenado" in m for m in r["motivos"])
    assert any("Binoculars" in a for a in r["avisos"])


class _BinocularsFalso:
    """Puntuación por palabra: las de la lista ``ia`` parecen IA (0.6) y el resto, humanas (1.2)."""
    model_id = "falso"

    def __init__(self, ia):
        self.ia = ia

    def token_stats(self, text):
        import re

        import numpy as np

        from backend.app.detectors.binoculars import TokenStats

        spans = [m.span() for m in re.finditer(r"\w+", text)]
        ppl = np.array([0.6 if text[a:b] in self.ia else 1.2 for a, b in spans])
        return TokenStats(offsets=np.array(spans), ppl=ppl, xppl=np.ones(len(spans)))


def test_titulos_no_se_marcan_por_encima_del_documento(monkeypatch):
    from backend.app import analyzer

    ia = " ".join(f"palabraia{i}" for i in range(35))
    texto = "Objetivo general\n\n" + ia + ".\n\n" + HUMANO + "\n\n" + HUMANO + "\n\n" + HUMANO
    monkeypatch.setattr(analyzer.binoculars, "get_detector", lambda: _BinocularsFalso(set(ia.split())))
    r = analyze_text(texto)
    assert r["probabilidad_ia"] < 0.35
    titulo, parrafo_ia = r["oraciones"][0], r["oraciones"][1]
    assert titulo["texto"] == "Objetivo general" and titulo["nivel"] == "bajo"
    assert parrafo_ia["nivel"] == "alto"  # un párrafo largo sí puede marcarse aunque el documento no
